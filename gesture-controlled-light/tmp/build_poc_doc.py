from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path(r'C:\Users\Kiran Rao\OneDrive\Documents\Gesture controlled light\gesture-controlled-light\Gesture-Controlled-Light-POC.docx')
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.62); sec.left_margin = Inches(.72); sec.right_margin = Inches(.72)
styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(9.5); styles['Normal'].font.color.rgb = RGBColor(45,55,68)
styles['Normal'].paragraph_format.space_after = Pt(5)
for nm,size,color in [('Title',29,'102A43'),('Heading 1',17,'0B7285'),('Heading 2',12,'155E75')]:
    st=styles[nm]; st.font.name='Aptos Display'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color); st.paragraph_format.space_before=Pt(9); st.paragraph_format.space_after=Pt(4)

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def set_cell_text(cell,text,bold=False,color='243B53',size=9):
    cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(1)
    r=p.add_run(text); r.bold=bold; r.font.name='Aptos'; r.font.size=Pt(size); r.font.color.rgb=RGBColor.from_string(color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Light Shading Accent 1'
    for i,h in enumerate(headers):
        set_cell_text(t.rows[0].cells[i],h,True,'FFFFFF',9); shade(t.rows[0].cells[i],'155E75')
    for row in rows:
        cells=t.add_row().cells
        for i,val in enumerate(row): set_cell_text(cells[i],val)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

def bullet(text):
    p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(2); p.add_run(text); return p

def code(text):
    t=doc.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    c=t.cell(0,0); shade(c,'F1F5F9'); c.text=''
    for idx,line in enumerate(text.strip('\n').splitlines()):
        p=c.paragraphs[0] if idx==0 else c.add_paragraph(); p.paragraph_format.space_after=Pt(0); p.paragraph_format.space_before=Pt(0)
        r=p.add_run(line or ' '); r.font.name='Consolas'; r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string('243B53')
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

def callout(label,text):
    t=doc.add_table(rows=1,cols=1); c=t.cell(0,0); shade(c,'E6F4F1'); c.text=''
    p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(2); r=p.add_run(label+'  '); r.bold=True; r.font.color.rgb=RGBColor.from_string('0B7285'); p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(8); r=p.add_run('PROOF OF CONCEPT  /  IoT + COMPUTER VISION'); r.bold=True; r.font.size=Pt(9); r.font.color.rgb=RGBColor.from_string('0B7285')
doc.add_paragraph('Gesture-Controlled\nPhone Light',style='Title')
p=doc.add_paragraph('A webcam hand-tracking prototype that controls an Android phone flashlight through nikki.build Streams.')
p.paragraph_format.space_after=Pt(12)
callout('PROJECT RESULT','Spread thumb and index finger to send flashlight ON; pinch them together to send OFF. Each successful command is followed by a vibration request.')

doc.add_heading('1. Purpose and scope', level=1)
doc.add_paragraph('This proof of concept uses a smartphone as the first IoT output device in a practical learning activity. It combines computer vision, gesture recognition, networked device control, and haptic feedback in one small, observable loop. The webcam and processing run on the computer; the phone supplies the flashlight and vibration hardware through an active Streams session.')
table(['Learning area','What students explore'],[
('Sensors and actuation','Camera input is interpreted as a control signal; phone flashlight and vibration act as outputs.'),
('Human-computer interaction','A natural hand motion becomes an explicit ON/OFF command with immediate feedback.'),
('Signal processing','Landmark distances are normalized, thresholded, stabilized across frames, and rate-limited.'),
('Networking and IoT','A Python SDK sends commands to a phone connected to a Streams session over the local network.'),
('Creative computing','Students can extend the same gesture events into games, accessibility controls, or other automations.')])

doc.add_heading('2. System design', level=1)
table(['Stage','Implementation'],[
('Capture','OpenCV reads webcam frames, mirrors the preview, and converts BGR to RGB.'),
('Perception','MediaPipe Hand Landmarker estimates 21 hand landmarks for one visible hand.'),
('Gesture logic','Thumb tip (landmark 4) to index tip (8) distance is divided by palm width (landmarks 5 and 17).'),
('Debounce','A gesture must remain stable for 4 frames; a 0.7-second minimum interval prevents rapid repeats.'),
('Actuation','Streams SDK sends set_light_status(True/False), then set_haptic() for a state change.'),
('Feedback','Camera overlay shows detected gesture and normalized ratio; terminal shows connection and command-send status.')])

# Deliberate second page for the technical details.
doc.add_page_break()
doc.add_heading('3. Gesture recognition details', level=1)
doc.add_paragraph('Raw pixel distance changes when the hand moves closer to or farther from the camera. Dividing fingertip distance by palm width gives a scale-normalized ratio. The current application defaults are PINCH_RATIO = 0.34, SPREAD_RATIO = 0.55, and STABLE_FRAMES = 4. Values between the thresholds are treated as neutral, so the controller does not force a state while the hand transitions.')
code('''def classify(hand, pinch_ratio=0.34, spread_ratio=0.55):
    thumb, index = hand[4], hand[8]
    palm = distance(hand[5], hand[17])
    ratio = distance(thumb, index) / max(palm, 1e-6)
    if ratio <= pinch_ratio:
        return "pinch", ratio       # flashlight OFF
    if ratio >= spread_ratio:
        return "spread", ratio      # flashlight ON
    return None, ratio               # neutral / transitioning''')
doc.add_paragraph('The production code also filters for a stable gesture across consecutive frames and only sends a command when the requested state differs from the last commanded state. This helps reduce flicker and duplicate network requests.')

doc.add_heading('4. Streams command path', level=1)
doc.add_paragraph('The session ID is requested through hidden terminal input and is not embedded in the project. The phone address is configured as host:port (without the ws:// prefix), matching the SDK interface used by this project.')
code('''SDK.initialize(session_id)
device = SDK.create_device("10.172.21.252:3000", on_event)
await device.connect()
await device.set_light_status(True)  # spread gesture
await device.set_haptic()            # vibration feedback
await device.set_light_status(False) # pinch gesture''')
callout('ACKNOWLEDGMENT LIMIT','The Streams SDK reports connection and send errors but does not provide a phone-side acknowledgment that a command was applied. The terminal message means the request was sent; verify the physical light/haptic on the phone.')

doc.add_heading('5. Project components', level=1)
table(['File / dependency','Role'],[
('main.py','Camera loop, MediaPipe landmark processing, gesture filter, reconnect handling, and Streams commands.'),
('models/hand_landmarker.task','Bundled hand landmark model used by MediaPipe; first launch does not need a model download.'),
('streams_smoke_test.py','Optional direct phone-control check independent of webcam and gesture classification.'),
('run.ps1 / requirements.txt','Windows launcher and pinned Python package versions.'),
('OpenCV / MediaPipe / nikki-streams','Video capture and display / hand landmarks / phone session communication.')])

doc.add_heading('6. Setup and operation', level=1)
for s in [
'Use a Windows computer with a webcam and Python 3.10+ (64-bit recommended). Install project dependencies from requirements.txt, or start from the project directory with: Set-ExecutionPolicy -Scope Process Bypass; .\\run.ps1.',
'Open the Streams app on the Android phone, join the intended session, and enable flashlight and haptics. Keep the phone and computer on the same local network.',
'Run python main.py. Enter the Streams Session ID at the hidden prompt. Show one well-lit hand; spread thumb and index to turn ON, pinch to turn OFF. Press Q or Esc in the preview to quit.',
'For an isolated connectivity check, run .venv\\Scripts\\python.exe .\\streams_smoke_test.py and observe the phone through the ON, wait, OFF sequence.'
]: bullet(s)

doc.add_heading('7. Proof-of-concept validation', level=1)
table(['Check','Observed outcome'],[
('Streams link and service access','Phone session joined; flashlight and haptics were enabled. Manual controls worked.'),
('Direct SDK test','User observed flashlight ON/OFF and vibration during the direct command sequence.'),
('Gesture-to-phone integration','User observed spread trigger ON plus vibration, and pinch trigger OFF plus vibration.'),
('Camera recovery','After opening the webcam shutter and allowing camera access, hand landmarks were detected in the preview.')])

doc.add_heading('8. Limitations and next steps', level=1)
for s in [
'Performance depends on lighting, camera framing, and keeping the full hand visible. The application provides an on-screen ratio and status to help tune placement.',
'The current controller supports one hand and two gestures. It does not yet persist a confirmed light state from the phone or query a phone-side acknowledgment.',
'For classroom extensions, students can add a visible state indicator, a calibration screen, gesture-controlled game events, event logging, or a second gesture for vibration patterns.',
'For deployment beyond a trusted local network, review session handling and transport security with the Streams platform documentation.'
]: bullet(s)

doc.add_heading('References', level=1)
for s in [
'Project source and usage notes: main.py, README.md, streams_smoke_test.py, requirements.txt.',
'MediaPipe Hand Landmarker for Python: https://ai.google.dev/edge/mediapipe/solutions/vision/hand_landmarker/python',
'nikki.build Streams Python SDK: https://github.com/nikki-build/streams.python'
]: bullet(s)

# Footer
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
rr=footer.add_run('GESTURE-CONTROLLED PHONE LIGHT  |  PROOF OF CONCEPT'); rr.font.name='Aptos'; rr.font.size=Pt(8); rr.font.color.rgb=RGBColor.from_string('829AB1')
doc.core_properties.title='Gesture-Controlled Phone Light Proof of Concept'
doc.core_properties.subject='Technical proof of concept, architecture, code snippets, setup, and validation'
doc.core_properties.author=''
doc.save(OUT)
print(OUT)
