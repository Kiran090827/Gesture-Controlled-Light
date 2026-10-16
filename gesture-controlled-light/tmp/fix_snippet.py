from docx import Document
from pathlib import Path
p=Path('Gesture-Controlled-Light-POC.docx')
d=Document(p)
for t in d.tables:
    if len(t.columns)==1:
        text='\n'.join(x.text for x in t.cell(0,0).paragraphs)
        if 'def classify(hand' in text:
            c=t.cell(0,0); c.text=''
            code='''import math

def classify(hand, pinch_ratio=0.34, spread_ratio=0.55):
    thumb, index = hand[4], hand[8]
    palm = math.hypot(hand[5].x - hand[17].x,
                      hand[5].y - hand[17].y)
    ratio = math.hypot(thumb.x - index.x,
                       thumb.y - index.y) / max(palm, 1e-6)
    if ratio <= pinch_ratio:
        return "pinch", ratio       # flashlight OFF
    if ratio >= spread_ratio:
        return "spread", ratio      # flashlight ON
    return None, ratio               # neutral / transitioning'''
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
            sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'F1F5F9'); c._tc.get_or_add_tcPr().append(sh)
            for i,line in enumerate(code.splitlines()):
                para=c.paragraphs[0] if i==0 else c.add_paragraph()
                para.paragraph_format.space_after=0
                r=para.add_run(line); r.font.name='Consolas'; r.font.size=__import__('docx').shared.Pt(8)
            break
d.save(p)
print(p.resolve())
