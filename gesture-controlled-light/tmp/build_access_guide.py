from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
OUT=Path(r'C:\Users\Kiran Rao\OneDrive\Documents\Gesture controlled light\gesture-controlled-light\Gesture-Controlled-Light-Project-Access-Guide.docx')
doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.78); sec.right_margin=Inches(.78)
st=doc.styles; st['Normal'].font.name='Aptos'; st['Normal'].font.size=Pt(10); st['Normal'].font.color.rgb=RGBColor.from_string('334E68'); st['Normal'].paragraph_format.space_after=Pt(6)
for n,z,c in [('Title',27,'102A43'),('Heading 1',16,'0B7285'),('Heading 2',12,'155E75')]:
 x=st[n]; x.font.name='Aptos Display'; x.font.size=Pt(z); x.font.bold=True; x.font.color.rgb=RGBColor.from_string(c); x.paragraph_format.space_before=Pt(9); x.paragraph_format.space_after=Pt(4)
def shade(cell,fill):
 pr=cell._tc.get_or_add_tcPr(); e=OxmlElement('w:shd'); e.set(qn('w:fill'),fill); pr.append(e)
def tab(headers,rows):
 t=doc.add_table(rows=1,cols=len(headers)); t.style='Light Shading Accent 1'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for i,h in enumerate(headers):
  c=t.rows[0].cells[i]; c.text=h; shade(c,'155E75'); r=c.paragraphs[0].runs[0]; r.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(9)
 for row in rows:
  cs=t.add_row().cells
  for i,v in enumerate(row):
   cs[i].text=v; cs[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   for r in cs[i].paragraphs[0].runs: r.font.size=Pt(9)
 doc.add_paragraph().paragraph_format.space_after=Pt(0)
def code(txt):
 t=doc.add_table(rows=1,cols=1); c=t.cell(0,0); shade(c,'F1F5F9'); c.text=''
 for i,line in enumerate(txt.strip().splitlines()):
  p=c.paragraphs[0] if i==0 else c.add_paragraph(); p.paragraph_format.space_after=Pt(0); r=p.add_run(line); r.font.name='Consolas'; r.font.size=Pt(9); r.font.color.rgb=RGBColor.from_string('243B53')
 doc.add_paragraph().paragraph_format.space_after=Pt(0)
def step(txt):
 p=doc.add_paragraph(style='List Number'); p.paragraph_format.space_after=Pt(4); p.add_run(txt)
p=doc.add_paragraph(); r=p.add_run('PROJECT HANDOFF  /  WINDOWS'); r.bold=True; r.font.size=Pt(9); r.font.color.rgb=RGBColor.from_string('0B7285')
doc.add_paragraph('Project Access Guide',style='Title')
doc.add_paragraph('How to locate, install, connect, and run the Gesture-Controlled Phone Light project.')
doc.add_heading('1. Get the project files',level=1)
doc.add_paragraph('The complete project folder is:')
code(r'C:\Users\Kiran Rao\OneDrive\Documents\Gesture controlled light\gesture-controlled-light')
doc.add_paragraph('A ZIP copy is available one folder above it:')
code(r'C:\Users\Kiran Rao\OneDrive\Documents\Gesture controlled light\gesture-controlled-light.zip')
doc.add_paragraph('To make a downloadable copy, open that folder in File Explorer and copy the ZIP to your preferred location, or right-click it and choose Extract All to unpack it. Keep the internal folder structure intact, including the models folder.')

doc.add_heading('2. What is included',level=1)
tab(['Item','Purpose'],[
('main.py','Main webcam, hand tracking, gesture, and phone-control program.'),
('models/hand_landmarker.task','Bundled MediaPipe hand model.'),
('requirements.txt','Pinned Python dependencies.'),
('run.ps1','Guided Windows launcher; creates/uses the virtual environment and installs dependencies.'),
('streams_smoke_test.py','Standalone test for Streams phone light and vibration.'),
('README.md','Project-level setup and usage notes.'),
('Gesture-Controlled-Light-POC.docx','Technical proof-of-concept report with architecture and code examples.'),
('Gesture-Controlled-Light-Documentation.pdf','Printable project documentation.')])

doc.add_heading('3. Computer setup and first launch',level=1)
for x in [
'Install 64-bit Python 3.10 or newer on the Windows computer. Python 3.12 is recommended for the documented setup.',
'Unzip the project if needed. Open PowerShell in the extracted gesture-controlled-light folder. In File Explorer, use the address bar to navigate to the folder, then type powershell and press Enter.',
'Run the launcher commands below. The first launch downloads/installs the pinned Python packages and can take a few minutes.',
'When prompted, enter the Streams Session ID from the phone app. The input is hidden while typing and should not be shared or hard-coded.',
'When the camera window appears, place one full hand in view. Spread thumb and index finger to request flashlight ON; pinch them together to request OFF. Press Q or Esc in the camera window to exit.'
]: step(x)
code('''Set-ExecutionPolicy -Scope Process Bypass
.\\run.ps1''')
doc.add_paragraph('Manual alternative, from the project directory:')
code('''py -3.12 -m venv .venv
.venv\\Scripts\\Activate.ps1
python -m pip install --upgrade pip
pip install -r requirements.txt
python main.py''')

doc.add_heading('4. Connect the phone with Streams',level=1)
for x in [
'Open nikki.build Streams on the Android phone and join the session intended for this project.',
'Confirm the flashlight and haptics services are enabled and permitted for the session. Keep the Streams app open while the Python program runs.',
'Ensure the phone and computer are on the same local Wi-Fi/LAN. The configured phone server address is 10.172.21.252:3000. The Python SDK expects host:port here, without the ws:// prefix.',
'Keep the Streams Session ID private. The project asks for it in hidden input; the Session ID is separate from the server address.'
]: step(x)

doc.add_heading('5. Verify connectivity independently',level=1)
doc.add_paragraph('If the camera or gesture detector is not the issue you are debugging, run the direct SDK test. Watch the phone: it sends ON plus vibration, waits about three seconds, then sends OFF plus vibration.')
code(r'.venv\Scripts\python.exe .\streams_smoke_test.py')
doc.add_paragraph('The test prompts for the same hidden Session ID. A terminal send message is not an application acknowledgment: the SDK does not report whether the phone physically applied each command.')

doc.add_heading('6. Troubleshooting',level=1)
tab(['Symptom','Try this'],[
('No camera / black frame','Close other camera apps; open the webcam shutter; allow desktop apps under Windows Settings > Privacy & security > Camera. Set CAMERA_INDEX to 1 if the desired webcam is not camera 0.'),
('Overlay says “Show one hand”','Move the hand back so the full hand is visible, face the palm toward the webcam, improve lighting, and avoid a busy background.'),
('Gesture detected but phone does not respond','Keep Streams open and joined; check session ID, phone server address, same-network access, and service permissions. Run streams_smoke_test.py to isolate the SDK connection.'),
('PowerShell blocks the launcher','Use the process-scoped command Set-ExecutionPolicy -Scope Process Bypass, then run .\\run.ps1. The change applies only to that PowerShell window.'),
('Light changes but vibration does not','Verify haptics are enabled in the Streams session. The main app sends haptics after a light state change; the standalone test sends a 500 ms request.')])

doc.add_heading('7. Gesture mapping and controls',level=1)
tab(['Gesture / key','Action'],[
('Pinch thumb and index fingertips','Send flashlight OFF, then request vibration.'),
('Spread thumb and index fingertips','Send flashlight ON, then request vibration.'),
('Q or Esc in the preview','Stop the application and disconnect Streams.')])
doc.add_paragraph('A hand-size-normalized distance and stability filter help reduce accidental commands. The flashlight state remains at its last commanded state when the program exits; check the phone manually if you need it left off.')

footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT; r=footer.add_run('GESTURE-CONTROLLED PHONE LIGHT  |  ACCESS GUIDE'); r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string('829AB1')
doc.core_properties.title='Gesture-Controlled Phone Light Project Access Guide'; doc.core_properties.author=''
doc.save(OUT); print(OUT)
